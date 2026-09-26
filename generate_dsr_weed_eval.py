"""
One-day DSR: in-crop weed model results and the validation weaknesses found.

Same layout as generate_dsr.py. Every figure below is a measured result from an
actual run. Nothing here reports a field test, because no field test has been
performed yet — that is itself one of the reported items.
"""

import os
import sys
import subprocess

try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REPORT_DATE = "2026-08-18"
FILENAME = "SmartKisan_DSR_2026-08-18.xlsx"

HEADERS = [
    "Date",
    "Task Category",
    "Detailed Task Description",
    "Deliverables / Artifacts",
    "Status",
    "Hours Spent",
    "Remarks & Notes",
]

ROWS = [
    [
        REPORT_DATE,
        "ML - In-Crop Model Training",
        "Completed the staged-transfer training for Green-on-Green detection. "
        "The backbone trained on DeepWeeds was carried over and fine-tuned on "
        "the in-crop CoFly patches, so the second stage only had to learn the "
        "harder problem of separating weed from crop when both are green, "
        "rather than learning weed appearance from 3,399 patches alone.",
        "cofly_model.tflite (2.4 MB), --init-from staged transfer",
        "Completed",
        "",
        "Validation accuracy 73.7%. Per-class recall: johnson_grass 91.3%, "
        "field_bindweed 90.0%, purslane 76.7%, background 69.2%. Lower than the "
        "rangeland model's 77.6%, which is expected - in-crop discrimination is "
        "the harder task and the earlier figure was measured on an easier one.",
    ],
    [
        REPORT_DATE,
        "ML - Defect Found (Evaluation)",
        "Identified that the reported in-crop accuracy is overstated by the "
        "evaluation method. The patch extractor tiles each source frame with "
        "50% overlap, producing roughly 17 patches per photograph, and the "
        "train/validation split then shuffles at patch level. Patches from the "
        "same photograph therefore appear on both sides of the split, so the "
        "model is being validated on scenes it trained on.",
        "Analysis of prepare_cofly.py splitting logic",
        "Open - Fix Identified",
        "",
        "Surfaced by field_bindweed reporting 90% recall from only 10 "
        "validation samples drawn from 14 source frames. Affects every class, "
        "not just that one, so 73.7% should be treated as optimistic. Fix is to "
        "split by source frame before tiling, so no photograph contributes to "
        "both sets, then retrain and re-measure.",
    ],
    [
        REPORT_DATE,
        "ML - Dataset Limitations",
        "Documented the coverage limits of the datasets currently in use. The "
        "in-crop set is 201 UAV frames from a single cotton field in Greece - "
        "one field, one crop, one season. The rangeland set used for the first "
        "stage photographs weeds against soil and scrub rather than inside a "
        "canopy. The crop-stress model is trained on single leaves captured "
        "under controlled lighting against uniform backgrounds.",
        "Model provenance notes, in-app disclosure",
        "Open - Data Required",
        "",
        "The crop-stress model's 99.2% should be read as a property of its "
        "dataset rather than as field accuracy; that figure is characteristic "
        "of controlled-condition imagery and is not expected to hold on field "
        "photographs. Moving to the in-crop model also dropped coverage of "
        "Parthenium and Lantana, both significant weeds in Indian conditions.",
    ],
    [
        REPORT_DATE,
        "ML - Validation Pending",
        "Field validation has not yet been carried out. The models are bundled "
        "in the build and every code path has been verified, but no photograph "
        "has been passed through them on a device, so no accuracy figure for "
        "real Indian crop images exists yet. Requires installing the build and "
        "scanning known weeds in the field.",
        "SmartKisan APK with on-device inference (71 MB)",
        "Blocked - Needs Device",
        "",
        "Priority check is the image preprocessing: the network expects inputs "
        "scaled to a specific range, and a mismatch there does not raise an "
        "error - it produces confident but meaningless predictions. Scanning a "
        "weed whose identity is already known will confirm or rule this out "
        "immediately. Until then the benchmark figures should not be quoted as "
        "field performance.",
    ],
]


def create_dsr():
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Status Report"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill("solid", start_color="2E7D32", end_color="2E7D32")
    title_font = Font(name="Segoe UI", size=16, bold=True, color="1B5E20")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="1E2C1F")
    bold_data_font = Font(name="Segoe UI", size=10, bold=True, color="1E2C1F")

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    title_align = Alignment(horizontal="left", vertical="center")

    thin = Side(style="thin", color="C8D6C8")
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A2"] = "SMARTKISAN PROJECT - DAILY STATUS REPORT (DSR)"
    ws["A2"].font = title_font
    ws["A2"].alignment = title_align

    ws["A3"] = (
        f"Weed detection model - training results and evaluation review, {REPORT_DATE}."
    )
    ws["A3"].font = Font(name="Segoe UI", size=10, italic=True, color="526654")
    ws["A3"].alignment = title_align

    for col, title in enumerate(HEADERS, 1):
        cell = ws.cell(row=5, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border_all
    ws.row_dimensions[5].height = 26

    row_num = 6
    for record in ROWS:
        for col, value in enumerate(record, 1):
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.font = bold_data_font if col == 5 else data_font
            cell.border = border_all
            cell.alignment = center if col in (1, 5, 6) else left
        ws.row_dimensions[row_num].height = 140
        row_num += 1

    for r in range(row_num, row_num + 10):
        ws.row_dimensions[r].height = 40
        for c in range(1, 8):
            cell = ws.cell(row=r, column=c)
            cell.font = data_font
            cell.border = border_all
            cell.alignment = center if c in (1, 5, 6) else left

    for col, width in {
        "A": 13, "B": 26, "C": 62, "D": 34, "E": 20, "F": 12, "G": 62
    }.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A6"

    path = os.path.join(os.getcwd(), FILENAME)
    wb.save(path)
    print(f"DSR written: {path}")


if __name__ == "__main__":
    create_dsr()

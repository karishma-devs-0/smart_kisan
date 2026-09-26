"""
Generates a one-day DSR covering bug fixes and defect resolution.

Same layout as generate_dsr.py so it drops into the existing reporting pack.
Every entry below corresponds to a real commit; nothing here is padding.
"""

import os
import sys
import subprocess

try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REPORT_DATE = "2026-08-11"
FILENAME = "SmartKisan_DSR_2026-08-11.xlsx"

HEADERS = [
    "Date",
    "Task Category",
    "Detailed Task Description",
    "Deliverables / Artifacts",
    "Status",
    "Hours Spent",
    "Remarks & Notes",
]

ROWS = [
    [
        REPORT_DATE,
        "Bug Fix - Authentication",
        "Fixed phone and username sign-in issuing an unusable session. Both "
        "methods returned a hardcoded placeholder token instead of a signed "
        "JWT, so the backend rejected it on every protected route: the user "
        "reached the dashboard apparently signed in while no data would load "
        "and nothing explained why. The same placeholder fallback was removed "
        "from email sign-in.",
        "src/services/api.js (authService)",
        "Fixed",
        "",
        "Latent until the app moved off mock data — once fields, crops, soil "
        "and profile started coming from the API, the broken session became "
        "visible as blanket load failures. Both methods now report that they "
        "are unavailable rather than granting a session that cannot work.",
    ],
    [
        REPORT_DATE,
        "Bug Fix - Account Deletion",
        "Fixed account deletion leaving user data behind. The delete routine "
        "removed only the row in the users table. With no cascade constraints "
        "on the user-scoped tables, the profile, fields, crops, devices, pumps "
        "and soil history all survived deletion indefinitely. Rewritten to "
        "clear every user-scoped table inside a single transaction.",
        "backend/src/routes/auth.js",
        "Fixed",
        "",
        "Compliance-relevant: Google Play requires an in-app deletion path, "
        "and deletion has to remove the data, not just the login. Transactional "
        "so a partial failure cannot leave an account half-deleted.",
    ],
    [
        REPORT_DATE,
        "Bug Fix - Notifications",
        "Fixed the notification feature never displaying anything. The screens, "
        "slice and storage helper were all present but nothing was connected: "
        "no code dispatched addNotification or setNotifications, the storage "
        "helper was never imported, and the socket client was never connected, "
        "so the server's real-time events had no listener. Wired MQTT alerts "
        "into the notification store and connected the socket on login.",
        "App.js, src/services/socketService.js",
        "Fixed",
        "",
        "The Home bell also rendered its unread dot unconditionally, showing an "
        "alert indicator whether or not anything was unread. It now reflects "
        "the real count and opens the notification list. Notification history "
        "persists across app restarts.",
    ],
    [
        REPORT_DATE,
        "Bug Fix - Real-Time / Config",
        "Fixed the socket client pointing at a service that does not exist. It "
        "targeted a hostname returning 404, used an /api path where Socket.IO "
        "does not attach, and opened a connection at import time before any "
        "user had signed in. Now derives the server origin from the configured "
        "API URL so the two cannot drift, and connects explicitly after login.",
        "src/services/socketService.js, src/services/backendApi.js",
        "Fixed",
        "",
        "Verified against the live deployment: Socket.IO handshake returns 200 "
        "and the existing REST routes continue to authenticate correctly.",
    ],
    [
        REPORT_DATE,
        "Issue Resolution - Data Integrity",
        "Removed code paths that presented generated values as real "
        "measurements. Disease detection returned a randomly chosen diagnosis "
        "with a confidence score and a chemical treatment whenever the model "
        "was unreachable, indistinguishable from a genuine result. Soil "
        "substituted placeholder moisture, pH and NPK readings when no sensor "
        "had reported. Both now report the failure or show an empty state.",
        "src/services/api.js, ScanResultScreen.js",
        "Fixed",
        "",
        "Highest-severity item of the set: acting on an invented diagnosis "
        "means applying the wrong chemical, and invented soil readings drive "
        "irrigation and fertiliser decisions. A scan that cannot run now fails "
        "with the reason instead of guessing.",
    ],
    [
        REPORT_DATE,
        "Bug Fix - Backend Stability",
        "Fixed the API rejecting all traffic behind the hosting proxy. The rate "
        "limiter refuses any request carrying an X-Forwarded-For header unless "
        "proxy trust is configured, which behind the platform proxy is every "
        "request. Configured trust for exactly one hop, so clients cannot spoof "
        "the header to bypass rate limiting.",
        "backend/src/server.js",
        "Fixed",
        "",
        "Would have failed on the first request after deployment. Verified with "
        "a simulated proxy header locally and with rapid repeat calls against "
        "the live service.",
    ],
]


def create_dsr():
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Status Report"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill("solid", start_color="2E7D32", end_color="2E7D32")
    title_font = Font(name="Segoe UI", size=16, bold=True, color="1B5E20")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="1E2C1F")
    bold_data_font = Font(name="Segoe UI", size=10, bold=True, color="1E2C1F")

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    title_align = Alignment(horizontal="left", vertical="center")

    thin = Side(style="thin", color="C8D6C8")
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A2"] = "SMARTKISAN PROJECT - DAILY STATUS REPORT (DSR)"
    ws["A2"].font = title_font
    ws["A2"].alignment = title_align

    ws["A3"] = (
        f"Defect resolution and stability fixes - {REPORT_DATE}. "
        "Authentication, account deletion, notifications, real-time connectivity "
        "and data integrity."
    )
    ws["A3"].font = Font(name="Segoe UI", size=10, italic=True, color="526654")
    ws["A3"].alignment = title_align

    for col, title in enumerate(HEADERS, 1):
        cell = ws.cell(row=5, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border_all
    ws.row_dimensions[5].height = 26

    row_num = 6
    for record in ROWS:
        for col, value in enumerate(record, 1):
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.font = bold_data_font if col == 5 else data_font
            cell.border = border_all
            cell.alignment = center if col in (1, 5, 6) else left
        ws.row_dimensions[row_num].height = 120
        row_num += 1

    for r in range(row_num, row_num + 10):
        ws.row_dimensions[r].height = 40
        for c in range(1, 8):
            cell = ws.cell(row=r, column=c)
            cell.font = data_font
            cell.border = border_all
            cell.alignment = center if c in (1, 5, 6) else left

    for col, width in {
        "A": 13, "B": 26, "C": 62, "D": 34, "E": 13, "F": 12, "G": 62
    }.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A6"

    path = os.path.join(os.getcwd(), FILENAME)
    wb.save(path)
    print(f"DSR written: {path}")


if __name__ == "__main__":
    create_dsr()

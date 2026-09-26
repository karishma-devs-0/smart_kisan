"""
One-day DSR: backend cold-start mitigation and code cleanup.

Same layout as generate_dsr.py. Figures below are measured, not estimated.
"""

import os
import sys
import subprocess

try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REPORT_DATE = "2026-08-26"
FILENAME = "SmartKisan_DSR_2026-08-26.xlsx"

HEADERS = ["Date", "Task Category", "Detailed Task Description",
           "Deliverables / Artifacts", "Status", "Hours Spent", "Remarks & Notes"]

ROWS = [
    [
        REPORT_DATE,
        "Backend - Cold Start Fix",
        "Diagnosed and mitigated the hosting platform's idle sleep. The free "
        "tier suspends the service after roughly 15 minutes without traffic, "
        "and the next request has to wait for it to restart. Nothing in the app "
        "woke the server before it was needed, so the whole restart landed on "
        "the user's first action - which was the sign-in itself.",
        "src/features/auth/screens/LoginScreen.js",
        "Fixed",
        "",
        "Measured 22.5s on a cold call against 0.26s warm. The login screen now "
        "pings the API when it opens, so the restart overlaps the seconds the "
        "user spends on the Google account picker. Fire-and-forget, so it never "
        "blocks the form, and it benefits email sign-in equally.",
    ],
    [
        REPORT_DATE,
        "Backend - Request Timeout",
        "Increased the API request budget and added a retry. The client allowed "
        "15 seconds per call, which is less than a cold start takes, so any "
        "request that arrived while the service was waking was abandoned "
        "mid-restart and surfaced as a connection failure.",
        "src/services/backendApi.js",
        "Fixed",
        "",
        "Now three attempts with an explicit 30s window each. HTTP errors are "
        "excluded from the retry - a wrong password should fail immediately "
        "rather than being retried against a verdict already received.",
    ],
    [
        REPORT_DATE,
        "ML Service - Cold Start Fix",
        "Applied the same treatment to the plant disease model, which is hosted "
        "separately and sleeps on the same basis. The scan screen now warms the "
        "model when it opens, so it is ready by the time the user has framed a "
        "photograph.",
        "src/services/api.js, DiseaseDetectionHomeScreen.js",
        "Fixed",
        "",
        "Measured 9.1s cold against 1.15s warm. Previously the timeout expired "
        "during the wake-up and the app fell back to placeholder output; that "
        "fallback has since been removed entirely.",
    ],
    [
        REPORT_DATE,
        "Code Cleanup - Dead Feature Code",
        "Removed the simulated camera-rig implementation from the field monitor "
        "screen after it was replaced with real on-device inference. The old "
        "implementation animated a scanning sequence and generated its results "
        "locally rather than from any model.",
        "WeedDetectionHomeScreen.js",
        "Completed",
        "",
        "1,420 lines reduced to 425. Removing it also made three dependencies "
        "unreachable - a graphics library, a slider control and a haptics "
        "module - all of which existed only for that screen.",
    ],
    [
        REPORT_DATE,
        "Code Cleanup - Dependencies & Package Size",
        "Removed the three orphaned dependencies and trimmed the build to the "
        "two processor architectures used by real phones. The package was "
        "shipping two further architectures that exist only for emulators.",
        "package.json, scripts/build-apk.sh",
        "Completed",
        "",
        "Package reduced from 134 MB to 71 MB. The architecture setting lives in "
        "a generated folder that is recreated whenever a native dependency "
        "changes, so an earlier fix had already been silently lost once; it is "
        "now passed by a build script that survives regeneration.",
    ],
    [
        REPORT_DATE,
        "Code Cleanup - Unused Imports & Strings",
        "Removed sample-data imports left behind when those screens moved onto "
        "the real API, and replaced the translation strings for the field "
        "monitor, which still described the removed simulation rather than the "
        "current screen.",
        "src/services/api.js, src/i18n/locales/en.js",
        "Completed",
        "",
        "Leaving unused sample-data imports in place makes it easy to "
        "reintroduce a fallback that silently replaces real data, which is what "
        "several earlier defects turned out to be.",
    ],
    [
        REPORT_DATE,
        "Build Cleanup - Upload Size",
        "Excluded machine-learning datasets, the backend and local documentation "
        "from the cloud build upload. The archive had reached 423 MB and the "
        "upload was failing partway through; application source is only a few MB "
        "of that.",
        ".easignore",
        "Completed",
        "",
        "The exclusion file replaces the normal ignore rules rather than adding "
        "to them, so it repeats the standard exclusions as well.",
    ],
]


def create():
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Status Report"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill("solid", start_color="2E7D32", end_color="2E7D32")
    title_font = Font(name="Segoe UI", size=16, bold=True, color="1B5E20")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="1E2C1F")
    bold_font = Font(name="Segoe UI", size=10, bold=True, color="1E2C1F")

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    title_align = Alignment(horizontal="left", vertical="center")

    thin = Side(style="thin", color="C8D6C8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A2"] = "SMARTKISAN PROJECT - DAILY STATUS REPORT (DSR)"
    ws["A2"].font = title_font
    ws["A2"].alignment = title_align

    ws["A3"] = ("Backend cold-start mitigation and code cleanup - " + REPORT_DATE + ".")
    ws["A3"].font = Font(name="Segoe UI", size=10, italic=True, color="526654")
    ws["A3"].alignment = title_align

    for col, title in enumerate(HEADERS, 1):
        c = ws.cell(row=5, column=col, value=title)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center
        c.border = border
    ws.row_dimensions[5].height = 26

    row = 6
    for record in ROWS:
        for col, value in enumerate(record, 1):
            c = ws.cell(row=row, column=col, value=value)
            c.font = bold_font if col == 5 else data_font
            c.border = border
            c.alignment = center if col in (1, 5, 6) else left
        ws.row_dimensions[row].height = 118
        row += 1

    for r in range(row, row + 10):
        ws.row_dimensions[r].height = 40
        for c in range(1, 8):
            cell = ws.cell(row=r, column=c)
            cell.font = data_font
            cell.border = border
            cell.alignment = center if c in (1, 5, 6) else left

    for col, w in {"A": 13, "B": 28, "C": 62, "D": 34, "E": 13,
                   "F": 12, "G": 62}.items():
        ws.column_dimensions[col].width = w

    ws.freeze_panes = "A6"

    path = os.path.join(os.getcwd(), FILENAME)
    wb.save(path)
    print("DSR written: " + path)
    print("  {} rows".format(len(ROWS)))


if __name__ == "__main__":
    create()

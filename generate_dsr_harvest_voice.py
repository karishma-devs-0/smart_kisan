"""
Three-day DSR: 28, 29 and 30 September 2026.

Deliberately short. The detail is in the feature list and the gap analysis;
this is the day-by-day record.

Written in the first person singular, like the other DSRs - it is submitted
under one name.
"""

import os
import sys
import subprocess

try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

MON = "2026-09-28"
TUE = "2026-09-29"
WED = "2026-09-30"
FILENAME = "SmartKisan_DSR_2026-09-28_to_09-30.xlsx"

HEADERS = ["Date", "Task", "What was done", "Status"]

ROWS = [
    # ── Monday ───────────────────────────────────────────────────────────
    [MON, "Harvest records",
     "Built the screen and the storage to record what actually comes off each "
     "field - quantity, area, quality and what was expected. Handles quintals, "
     "kilograms, tonnes and maunds over acres, bigha or hectares.",
     "Completed"],
    [MON, "Yield prediction",
     "Now works, because harvests are recorded. It shows the farmer's own "
     "average yield per acre over the area planted. It was showing the same "
     "invented figures to everyone before.",
     "Completed"],
    [MON, "Harvest report",
     "The harvest performance report shows real figures. It used to show 91.7% "
     "efficiency to every farmer.",
     "Completed"],

    # ── Tuesday ──────────────────────────────────────────────────────────
    [TUE, "Weather forecast",
     "Days six to fourteen of the forecast were being made up - the first five "
     "days repeated with random changes. They now come from a service that "
     "actually forecasts fourteen days.",
     "Fixed"],
    [TUE, "Weather advice",
     "Added what the forecast means for the week's work: which days are fit to "
     "spray and why not on the others, plus heat, frost and heavy rain "
     "warnings with the day named.",
     "Completed"],
    [TUE, "Weed detection",
     "Stopped the model guessing whether a plant is the crop - it gets that "
     "right only 23 times in 100. The app answers it from the crop the farmer "
     "entered instead. Worth about three points.",
     "Fixed"],
    [TUE, "Report screens",
     "Three report screens had their figures written into the code, so the "
     "real data from the server never reached the farmer. Connected them "
     "properly.",
     "Fixed"],

    # ── Wednesday ────────────────────────────────────────────────────────
    [WED, "App reads aloud",
     "The app now reads its advice out in the farmer's own language - the "
     "disease result, the weed result, farm observations and weather warnings. "
     "This matters for farmers who cannot read comfortably.",
     "Completed"],
    [WED, "App listens",
     "The farmer can speak to move around the app, in Hindi, Punjabi and "
     "English. It only navigates and never saves anything, so a misheard word "
     "cannot get written into a record.",
     "Completed"],
    [WED, "Hindi and Punjabi",
     "Translated the new screens into Hindi and Punjabi rather than leaving "
     "them in English.",
     "Completed"],
    [WED, "Testing and build",
     "Built a new APK and checked the voice commands with an automated test - "
     "43 spoken phrases across all three languages. Updated the testing "
     "manual, the feature list and the gap analysis.",
     "Completed"],
]

NOTES = [
    "A new APK is ready for the testing team. The voice features need the "
    "installed APK - they cannot run in Expo Go.",
    "The microphone works in Hindi, Punjabi and English only. The phone's "
    "recogniser handles the other seven languages poorly, so it is hidden "
    "there rather than offered and disappointing.",
    "The Hindi and Punjabi wording should be read by a native speaker before "
    "release.",
]


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "DSR"

    ws.append(["SmartKisan - Daily Status Report"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    c = ws.cell(row=1, column=1)
    c.font = Font(bold=True, size=14, color="1B5E20")
    c.alignment = Alignment(horizontal="center")

    ws.append(["28 - 30 September 2026"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
    c = ws.cell(row=2, column=1)
    c.font = Font(italic=True, size=10, color="555555")
    c.alignment = Alignment(horizontal="center")

    ws.append([])
    ws.append(HEADERS)

    thin = Side(style="thin", color="BDBDBD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for i in range(1, 5):
        cell = ws.cell(row=4, column=i)
        cell.fill = PatternFill("solid", fgColor="2E7D32")
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    colours = {"Completed": "1B5E20", "Fixed": "1B5E20"}

    for row in ROWS:
        ws.append(row)
        r = ws.max_row
        for i in range(1, 5):
            cell = ws.cell(row=r, column=i)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=r, column=4).font = Font(
            bold=True, color=colours.get(row[3], "000000"))

    ws.append([])
    ws.append(["Notes"])
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.cell(row=r, column=1).font = Font(bold=True, size=12, color="1B5E20")
    ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor="C8E6C9")

    for note in NOTES:
        ws.append(["", note])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.cell(row=r, column=2).alignment = Alignment(vertical="top", wrap_text=True)

    for i, w in enumerate([12, 24, 78, 14], start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    ws.row_dimensions[4].height = 22
    ws.freeze_panes = "A5"

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), FILENAME)
    wb.save(out)
    print("Wrote " + out)
    print("  %d rows over 3 days" % len(ROWS))


if __name__ == "__main__":
    build()

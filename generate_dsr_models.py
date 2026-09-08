"""
DSR for Thursday 3 September 2026.

Covers the work after the 31 Aug - 2 Sep report: the evaluation harnesses
for both models, and the disease model fine-tune.

Same layout as the previous DSR generators. Every row corresponds to a real
commit on main.
"""

import os
import sys
import subprocess

try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

THU = "2026-09-03"
FILENAME = "SmartKisan_DSR_2026-09-03.xlsx"

HEADERS = ["Date", "Task Category", "Detailed Task Description",
           "Deliverables / Artifacts", "Status", "Hours Spent", "Remarks & Notes"]

ROWS = [
    [
        THU,
        "AI Models - Measurement",
        "Built test harnesses for both AI features. Until now the only figure "
        "either model had was its training accuracy, which is measured on "
        "images from the same collection it learned from and reports "
        "memorisation rather than usefulness. Both are now scored against "
        "photographs taken elsewhere, in real conditions.",
        "evaluate_weed_model.py, evaluate_disease_model.py, "
        "evaluate_field_model.py",
        "Completed",
        "",
        "This is the groundwork for every result below. Neither model could "
        "previously be judged before release.",
    ],
    [
        THU,
        "Disease Detection - Field Accuracy",
        "The disease model had only ever seen laboratory photographs: one "
        "detached leaf on a plain background under even light. The app only "
        "ever receives photographs taken in a field. Retrained it on field "
        "imagery while retaining what it already knew.",
        "finetune_field.py, field_model.tflite",
        "Ready to Deploy",
        "",
        "Correct diagnoses rose from 28% to 58% on real field photographs. "
        "More importantly, the old model called 47% of HEALTHY plants "
        "diseased, recommending a fungicide that was not needed; that is now "
        "8%. Laboratory accuracy was barely affected, 99% to 97%.",
    ],
    [
        THU,
        "Disease Detection - Safety Threshold",
        "Raised the confidence level below which the app declines to name a "
        "disease, from 60% to 70%, based on measurement across the same test "
        "set. At 60% the app answers 56% of scans with 75% of those correct; "
        "at 70% it answers 47% with 84% correct.",
        "src/services/api.js",
        "Fixed",
        "",
        "Nearly halves the wrong diagnoses a farmer is shown. An unanswered "
        "scan costs someone another photograph; a wrong one costs them a "
        "fungicide and a sprayed plant that did not need it.",
    ],
    [
        THU,
        "Weed Detection - Training Experiments",
        "Tested whether adding two further public datasets would improve weed "
        "identification, as requested. Both were trained and measured. Both "
        "made the model worse on real photographs and neither is used.",
        "train_weed_model.py (--task combined)",
        "Investigated - Not Adopted",
        "",
        "Accuracy on unfamiliar photographs fell from 87% to 73% with one "
        "dataset added and to 53% with both. The added imagery was Australian "
        "rangeland and drone footage - neither resembles a farmer's phone "
        "photograph. The existing model remains in the app and is the better "
        "one.",
    ],
    [
        THU,
        "Weed Detection - Test Coverage",
        "The weed model was being judged on 15 photographs, so a difference of "
        "two images looked like a meaningful result. Added two further test "
        "sets already available on disk, giving 1,284 images the model has "
        "never seen.",
        "evaluate_weed_model.py",
        "Completed",
        "",
        "The earlier conclusion held on the larger sample and the margin "
        "widened. It also revealed a weakness the small set could not: grass "
        "weeds are identified correctly only 13% of the time in unfamiliar "
        "conditions. Grass versus broadleaf decides which herbicide to buy, so "
        "this is the priority for the next round of work.",
    ],
    [
        THU,
        "Weed Detection - Incorrect Information Shown",
        "The note displayed to the farmer described a model that is not in the "
        "app, naming datasets replaced in August. Corrected to describe what "
        "actually runs, including that grass weeds are its weakest class.",
        "WeedDetectionHomeScreen.js, weedInference.js",
        "Fixed",
        "",
        "Telling someone the wrong thing about how a spraying decision was "
        "reached is worse than telling them nothing.",
    ],
    [
        THU,
        "Model Tuning - Three Further Attempts",
        "Tried three further techniques to improve the disease model beyond "
        "58%: a larger training run, averaging over multiple views of each "
        "photograph, and weighting rare diseases more heavily. All three were "
        "measured and none is adopted.",
        "finetune_field.py",
        "Investigated - Not Adopted",
        "",
        "Each result is recorded in the code beside the setting it concerns so "
        "the work is not repeated. The conclusion is that the remaining "
        "weakness is a shortage of training photographs for specific diseases, "
        "not a shortage of tuning.",
    ],
]

OPEN_ITEMS = [
    ["Disease model", "Trained and verified, not yet live",
     "Needs a HuggingFace access token to upload. Deployment script is written "
     "and tested; the upload itself takes a few minutes once the token exists."],
    ["Phone sign-in", "Codes cannot reach Indian numbers",
     "Requires DLT registration with TRAI, or credentials for an SMS account "
     "the company already holds. Business registration, not development."],
    ["TC-016", "Disease detection reported as showing a rejection",
     "The service was tested directly and responds correctly. Awaiting the "
     "exact error text from the testing team."],
    ["TC-008", "Username login",
     "The app identifies accounts by email or phone; there is no username "
     "system. Replaced with email code sign-in. Awaiting confirmation that "
     "this is acceptable."],
    ["Both AI models", "Further accuracy is limited by photographs, not code",
     "What would help most is real photographs from Indian fields - wheat, "
     "rice and cotton for weeds, and specific tomato diseases. Around 50 per "
     "category would be worth more than further tuning."],
    ["Real-time messaging", "Pump control uses a public broker",
     "Anyone who knows a user ID could read sensor data or switch a pump. "
     "Needs a private broker before physical hardware is connected."],
]

HEADER_FILL = PatternFill("solid", fgColor="2E7D32")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SECTION_FILL = PatternFill("solid", fgColor="C8E6C9")
THIN = Side(style="thin", color="BDBDBD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_COLOURS = {
    "Fixed": "1B5E20",
    "Completed": "1B5E20",
    "Ready to Deploy": "0277BD",
    "Investigated - Not Adopted": "E65100",
}


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "DSR"

    ws.append(["SmartKisan - Daily Status Report"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.cell(row=1, column=1).font = Font(bold=True, size=14, color="1B5E20")
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    ws.append(["Thursday, 3 September 2026 (following the 31 Aug - 2 Sep report)"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="center")
    ws.cell(row=2, column=1).font = Font(italic=True, size=10, color="555555")

    ws.append([])
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=4, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for row in ROWS:
        ws.append(row)
        r = ws.max_row
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        colour = STATUS_COLOURS.get(row[4])
        if colour:
            ws.cell(row=r, column=5).font = Font(bold=True, color=colour)

    ws.append([])
    ws.append(["Open Items"])
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEADERS))
    ws.cell(row=r, column=1).font = Font(bold=True, size=12, color="1B5E20")
    ws.cell(row=r, column=1).fill = SECTION_FILL

    ws.append(["Reference", "Item", "Blocked On / Next Step", "", "", "", ""])
    r = ws.max_row
    for c in range(1, 4):
        cell = ws.cell(row=r, column=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8F5E9")
        cell.border = BORDER

    for item in OPEN_ITEMS:
        ws.append([item[0], item[1], item[2], "", "", "", ""])
        r = ws.max_row
        for c in range(1, 4):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for i, w in enumerate([12, 30, 62, 32, 24, 12, 48], start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    ws.row_dimensions[4].height = 28
    ws.freeze_panes = "A5"

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), FILENAME)
    wb.save(out)
    print("Wrote " + out)
    print("  %d task rows, %d open items" % (len(ROWS), len(OPEN_ITEMS)))


if __name__ == "__main__":
    build()

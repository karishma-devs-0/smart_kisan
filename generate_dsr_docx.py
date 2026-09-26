"""
Builds the Thursday-to-Saturday status report as a Word document.

Plain language throughout: this goes to management, so it describes what was
wrong and what changed rather than naming files or functions.
"""

import os
import subprocess
import sys

try:
    import docx  # noqa: F401
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

FILENAME = "SmartKisan_Status_Report_27-29_Aug_2026.docx"

GREEN = RGBColor(0x1B, 0x5E, 0x20)
GREY = RGBColor(0x52, 0x66, 0x54)

DAYS = [
    (
        "Thursday, 27 August 2026",
        "QA Defect Resolution",
        "Reviewed the testing team's report and resolved the issues raised.",
        [
            ("Weed scan produced no result",
             "The photograph was captured and analysed, but the result was read back from "
             "the analysis engine in the wrong format, producing an empty answer instead of "
             "an error. The screen therefore had nothing to display. The model itself was "
             "never at fault — confirmed afterwards by running the application's exact "
             "processing sequence over the reference images, with 12 of 12 correct."),
            ("Pumps and soil readings were not being saved",
             "Both features were still using sample data. Adding a pump created a temporary "
             "entry that never reached the server, so it appeared in the list and was gone "
             "after restarting the application; adding a soil reading did nothing at all. "
             "Every account also saw the same sample pumps. Reported as three separate "
             "items, with one underlying cause."),
            ("Registration accepted invalid details",
             "Only the presence of a value was checked, so an entry such as 'abc' was "
             "accepted as an email address and any sequence of digits as a phone number. "
             "Format checks were added, with the phone rule matching Indian mobile numbers."),
            ("Sign-in error remained on screen after switching method",
             "An error from one sign-in method stayed visible after moving to another, so a "
             "failed code was still displayed while the user was on the email form."),
            ("Weather showed the wrong location",
             "The farm's location is saved with the account, but nothing restored it after a "
             "reinstall or on a second device, so a default city was used instead."),
            ("Pump form always read 'Edit'",
             "A single screen serves both adding and editing a pump, and always showed the "
             "edit heading."),
        ],
    ),
    (
        "Friday, 28 August 2026",
        "Backend Reliability and Verification",
        "Restored the service after a hosting suspension, then addressed reliability gaps "
        "found while reviewing the server.",
        [
            ("Hosting service was suspended",
             "The backend became unavailable. A new service was created and the API "
             "redeployed to restore it, with configuration re-entered and the deployment "
             "verified against the live endpoints."),
            ("Health check reported healthy without checking anything",
             "The hosting platform uses this to decide whether the service is working, so an "
             "instance that had lost its database stayed in service while every request "
             "failed — the one situation the check exists to catch. It now queries the "
             "database and reports unhealthy with the reason."),
            ("No orderly shutdown",
             "Requests in progress were cut off during deployments, and database connections "
             "were abandoned rather than closed. The server now finishes current work before "
             "stopping."),
            ("Login attempts were barely limited",
             "Sign-in allowed 100 requests a minute, which is 100 password guesses a minute. "
             "Now limited to 10 attempts per 15 minutes, counted per account rather than per "
             "network address, so attempts against one account cannot lock out other users "
             "sharing a mobile network."),
            ("Submitted data went unchecked into the database",
             "A typing mistake produced a generic server error that told the user nothing and "
             "filled the logs with database errors that were really client mistakes. Invalid "
             "input now returns a message naming the field."),
            ("Automated verification built",
             "A single command now exercises the whole user journey — sign up, sign in, farm "
             "setup, records, sensor readings and account deletion — reporting pass or fail "
             "per step. All 40 checks passing. This also tested account deletion for the "
             "first time; the code existed but had never been run."),
        ],
    ),
    (
        "Saturday, 29 August 2026",
        "Sign-in Options and Stability",
        "Added alternative ways to sign in and fixed two faults that prevented the "
        "application from starting.",
        [
            ("Sign-in by email code added",
             "Along with address verification and password reset. Password reset closes a "
             "real gap: there was previously no recovery path at all, so a forgotten "
             "password meant a permanently lost account."),
            ("Two broken sign-in tabs replaced",
             "Phone and Username could never work — there is no messaging provider and no "
             "username lookup — so both existed only to report themselves unavailable, which "
             "the testing team reasonably recorded as broken."),
            ("Application crashed at startup in the development preview",
             "A component loaded the on-device detection engine during startup, and the "
             "preview environment cannot provide it, so the whole application failed before "
             "rendering rather than only that one feature."),
            ("Blank white screen after opening the application",
             "The application waited for a server response before deciding which screen to "
             "show, and displayed nothing at all while waiting. Because the server sleeps "
             "when idle, that wait can run to most of a minute, which is indistinguishable "
             "from a crash."),
            ("Farm location now requested during setup",
             "Only a district was collected previously, so the farm map showed a point in the "
             "middle of a city rather than the farmer's own land. The farmer now marks their "
             "field on a map."),
        ],
    ),
]


def add_heading(doc, text, size, color, space_before=0, space_after=6, bold=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = "Segoe UI"
    return p


def build():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Segoe UI"
    style.font.size = Pt(10)

    for section in doc.sections:
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)

    add_heading(doc, "SMARTKISAN", 20, GREEN, space_after=0)
    add_heading(doc, "Status Report — 27 to 29 August 2026", 13, GREY,
                space_after=4, bold=False)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(14)
    r = sub.add_run("Defects resolved across the mobile application and backend service.")
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = GREY

    total = 0
    for date_label, title, intro, items in DAYS:
        add_heading(doc, date_label, 12, GREEN, space_before=12, space_after=0)
        add_heading(doc, title, 10.5, GREY, space_after=4, bold=False)

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(intro)
        r.font.size = Pt(9.5)

        for name, detail in items:
            total += 1
            b = doc.add_paragraph(style="List Bullet")
            b.paragraph_format.space_after = Pt(6)
            head = b.add_run(name + ". ")
            head.bold = True
            head.font.size = Pt(9.5)
            body = b.add_run(detail)
            body.font.size = Pt(9.5)

    add_heading(doc, "Summary", 12, GREEN, space_before=16, space_after=4)
    s = doc.add_paragraph()
    r = s.add_run(
        f"{total} issues resolved over three days, covering data not being saved, "
        "invalid input being accepted, two faults that prevented the application from "
        "starting, a hosting suspension, and gaps in server reliability and security. "
        "Automated verification of the backend was introduced and is passing in full."
    )
    r.font.size = Pt(9.5)

    path = os.path.join(os.getcwd(), FILENAME)
    doc.save(path)
    print("Report written: " + path)
    print(f"  {len(DAYS)} days, {total} items")


if __name__ == "__main__":
    build()

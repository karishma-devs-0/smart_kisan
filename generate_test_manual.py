"""
User manual for testing the SmartKisan APK.

Written for the testing team: every screen in the app, what to do on it, and
what should happen. Where something is known not to work, it says so and why,
so a known limit is not filed as a defect and time is not wasted on it.

SCREENSHOTS
-----------
Each step has a place for one. If docs/screenshots/<name>.png exists it is
placed in the document; if not, a labelled box appears saying which screenshot
belongs there. Capture them with:

    bash scripts/capture-screens.sh      (phone plugged in, USB debugging on)
    python generate_test_manual.py       (again, to place them)

A phone rather than an emulator, because the camera and the on-device weed
detector do not run in one, and a screenshot from an emulator looks nothing
like what a farmer sees.
"""

import os
import subprocess
import sys

try:
    import docx  # noqa: F401
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, "docs", "screenshots")
FILENAME = "SmartKisan_Testing_Manual.docx"

GREEN = RGBColor(0x1B, 0x5E, 0x20)
AMBER = RGBColor(0xE6, 0x51, 0x00)
GREY = RGBColor(0x77, 0x77, 0x77)
DARK = RGBColor(0x33, 0x33, 0x33)

# (shot, title, steps[], expected, note or None)
# `note` is a known limitation - something that will look wrong but is not.
SECTIONS = [
    ("Getting started", [
        ("01-install", "Install the app", [
            "Copy app-release.apk to the phone.",
            "Open it. Android will warn about installing outside the Play Store - "
            "allow it for this once.",
            "Open SmartKisan from the app drawer.",
        ], "The app opens on the consent screen.",
         "Android 13 and newer ask for notification permission on first open. "
         "Allow it, or the alert screens cannot be tested."),

        ("02-consent", "Terms and consent", [
            "Read the terms.",
            "Scroll to the bottom.",
            "Tick the box and continue.",
        ], "The tick box can be ticked as soon as you reach the bottom, and "
           "Continue then works.",
         "This used to need scrolling up and down again before the box would "
         "respond. If that happens, it is a defect - please report it."),
    ]),

    ("Creating an account", [
        ("03-register", "Register", [
            "Tap Register.",
            "Enter a name, an email you can receive post at, a 10-digit mobile "
            "number and a password of at least 6 characters.",
            "Submit.",
        ], "The account is created and you land on farm setup.",
         None),

        ("04-register-error", "Registration refuses bad details", [
            "Try an email without an @.",
            "Try a mobile number of 5 digits.",
            "Try registering again with an email already used.",
        ], "Each is refused with a message under the field saying what is "
           "wrong. The email already in use says so specifically.",
         None),

        ("05-login-email", "Sign in with a password", [
            "Sign out, then sign in with the email and password.",
        ], "You reach the dashboard, and your farm is there - the farm name, "
           "fields and crops you entered during setup.",
         "The first sign-in of the day can take up to a minute. The server "
           "sleeps when unused and has to wake. A spinner shows throughout. "
           "This is the hosting plan, not a fault."),

        ("06-login-code", "Sign in without a password", [
            "On the login screen, choose the Code tab.",
            "Enter your email and ask for a code.",
            "Check your email.",
        ], "A 6-digit code arrives. Entering it signs you in.",
         "The code lasts 10 minutes, allows 5 attempts and works once. "
           "Asking for a new code cancels the previous one - so if two arrive, "
           "only the newer works."),

        ("07-code-entered", "The code screen", [
            "Look at the wording after the code is sent.",
        ], "It says a code is on its way if an account exists for that email.",
         "It says 'if an account exists' deliberately, whether or not one "
           "does. Confirming which emails are registered would let anyone "
           "test for our users' addresses."),

        ("08-forgot-password", "Reset a forgotten password", [
            "On the Email tab, tap 'Forgot password?'.",
            "Enter your email and request a code.",
            "Enter the code and a new password.",
        ], "The password changes and you are signed straight in. The old "
           "password no longer works.",
         None),
    ]),

    ("Setting up the farm", [
        ("09-onboard-farm", "Farm name and type", [
            "Enter a farm name and choose the type.",
        ], "Both are accepted and Next works.", None),

        ("10-onboard-size", "Farm size", [
            "Choose a size band, or enter the exact area.",
        ], "Either is accepted.", None),

        ("11-onboard-location", "Where the farm is", [
            "Search for your town or village and select it.",
        ], "The location is set. This is what the weather screens and the farm "
           "map use later.",
         "Worth setting accurately. If weather later shows the wrong place, "
           "this is the screen to check first."),

        ("12-onboard-fields", "Fields, crops and devices", [
            "Add at least two fields with areas.",
            "Add crops and assign them to fields.",
            "Add a device if you have one.",
        ], "Setup finishes and the dashboard opens. Everything entered here is "
           "saved on the server, not just on the phone.",
         "Reinstall the app and sign in again to check - the farm should "
           "still be there."),
    ]),

    ("The dashboard", [
        ("13-home", "Home", [
            "Look at the top cards.",
        ], "Pumps running, soil moisture and temperature. A dash means no "
           "reading yet, which is correct - not a fault.",
         None),

        ("14-home-summary", "Today's water and electricity", [
            "Scroll to Today's Summary.",
            "Run a pump for a few minutes, come back and pull to refresh.",
        ], "Hours, litres and units change to reflect the run. Litres come "
           "from the pump's flow rate, units from its horsepower.",
         "Before any pump has run today it shows dashes, not zeros. These "
           "used to be the same numbers at every sign-in; if you see that "
           "again, report it."),

        ("15-farm-map", "Farm map", [
            "Look at the map on the dashboard, then open the full map.",
        ], "Your own fields are drawn on satellite imagery.", None),

        ("16-notifications", "Notifications", [
            "Tap the bell.",
        ], "The alert list opens. The red dot appears only when something is "
           "unread.",
         None),
    ]),

    ("Pumps", [
        ("17-pumps-list", "Pump list", [
            "Open the Pump tab.",
        ], "Every pump, by category, with its state.", None),

        ("18-pump-add", "Add a pump", [
            "Add a pump with a name, type, flow rate in litres per minute and "
            "horsepower.",
            "Save it.",
            "Close the app completely and reopen it.",
        ], "The pump is still there after reopening.",
         "This is worth testing carefully. Adding a pump used to appear to "
           "work and the pump was gone after a restart, because it was never "
           "saved. Flow rate and horsepower matter - the water and "
           "electricity figures are worked out from them."),

        ("19-pump-detail", "Pump detail", [
            "Open a pump.",
        ], "Its details, state and recent activity.", None),

        ("20-pump-control", "Switch a pump on and off", [
            "Turn a pump on. Wait a minute. Turn it off.",
        ], "It switches both ways with no error. The state on the list "
           "updates without refreshing.",
         "Turning a pump on used to show an error even though the pump had "
           "switched. If you see an error here, check whether the pump "
           "actually changed state before reporting it."),

        ("21-pump-timer", "Timer", [
            "Set a pump to run for 2 minutes.",
        ], "It counts down and stops on its own.", None),

        ("22-pump-history", "Pump history", [
            "Open history after running a pump.",
        ], "The run is listed with its duration.",
         "No run was recorded at all before this release, so history was "
           "always empty. It should not be now."),
    ]),

    ("Soil", [
        ("23-soil", "Soil readings", [
            "Open the Soil tab.",
        ], "Moisture, pH, nitrogen, phosphorus, potassium and organic carbon.",
         "With no sensor and no reading entered, these are blank. Blank is "
           "correct - a made-up 45% would look like a real measurement."),

        ("24-soil-add", "Add a reading by hand", [
            "Add a soil reading with values from a test kit or lab report.",
            "Close the app and reopen it.",
        ], "The reading is saved and still there after reopening.", None),

        ("25-soil-health", "Soil health", [
            "Open the soil health screen.",
        ], "A score and what is short. Enter a low nitrogen value, like 200, "
           "and it should say nitrogen is low.",
         "Ratings follow Soil Health Card bands. The nutrient rating takes "
           "the worst of N, P and K, not the average - so one low value pulls "
           "the whole rating down. That is deliberate."),
    ]),

    ("Weather", [
        ("26-weather-today", "Today", [
            "Open the Sky tab.",
        ], "Temperature, humidity and conditions for your farm's location.",
         "If this shows the wrong place, check the farm location in Settings."),

        ("27-weather-forecast", "Forecast", ["Open the forecast."],
         "The days ahead.", None),

        ("28-weather-et", "Water need calculator", [
            "Open the ET calculator and pick a crop.",
        ], "How much water the crop needs.", None),
    ]),

    ("Crops, fields and devices", [
        ("29-crops", "My crops", ["Open My Crops."],
         "Crops with their field and stage.", None),
        ("30-crop-add", "Add a crop", [
            "Add a crop, assign it to a field, set a sowing date.",
            "Reopen the app.",
        ], "It is still there.", None),
        ("31-fields", "My fields", ["Open My Fields."],
         "Each field with its area and soil type.", None),
        ("32-devices", "Devices", ["Open the device list."],
         "Devices with their online state and battery.",
         "No physical sensors have been built yet, so this is tested with "
           "simulated devices. An empty list is expected on a new account."),
    ]),

    ("The AI features", [
        ("33-disease-home", "Disease detection", [
            "Settings, then Disease Detection.",
        ], "The scan screen, offering camera or gallery.",
         "Give it a few seconds on first open. The model service sleeps when "
           "unused and the app wakes it in the background."),

        ("34-disease-result", "Scan a leaf", [
            "Photograph a single leaf, filling the frame, in daylight, "
            "against a plain background.",
        ], "The crop, the disease, a treatment with dosage, and organic "
           "alternatives.",
         "Covers 9 crops: tomato, potato, maize, rice, wheat, capsicum, "
           "soybean, squash and orange. A leaf from another crop will be "
           "matched to the nearest of these and be wrong - that is a limit of "
           "the model, not a defect."),

        ("35-disease-refused", "When it will not answer", [
            "Photograph something that is not a leaf - a hand, or the ground.",
        ], "It says 'Try another photo' and explains what to change.",
         "This is the safety check working, not a failure. It refuses below "
           "80% confidence rather than naming a disease from a guess. It also "
           "refuses to call a plant healthy unless it is confident - telling "
           "a farmer a sick plant is fine is the costliest mistake it can "
           "make. Please do not file this as a defect."),

        ("36-weed-home", "AI field monitor", [
            "Settings, then AI Field Monitor.",
        ], "The scan screen with two options, green-on-green and "
           "yellow-on-green.",
         "This only works on the installed APK. It cannot run in Expo Go at "
           "all - it is a native component. If it was tested in Expo Go "
           "before and did nothing, that is why."),

        ("37-weed-result", "Scan a plant", [
            "Photograph a weed or a crop plant, 20-40 cm away.",
        ], "Grass weed, broadleaf weed, or crop.",
         "93% accurate at telling grass from broadleaf on real photographs. "
           "It is less reliable at deciding whether a plant is the crop - it "
           "learned that from one crop on one farm."),
    ]),

    ("Advice and planning", [
        ("38-crop-suitability", "Crop suitability", [
            "Settings, then Crop Suitability.",
        ], "Crops suited to your soil, with reasons.",
         "With no soil reading it shows nothing and asks for one. That is "
           "correct - recommending a crop from invented soil figures would be "
           "worse than recommending nothing. Values can also be typed in."),

        ("39-fertilizer", "Fertiliser calculator", [
            "Pick a crop, a target yield and a field area.",
        ], "How much N, P and K, and what it costs.", None),

        ("40-tasks", "Farm tasks", ["Open Farm Management, then tasks."],
         "Your own task list. Empty on a new account.",
         "This used to show the same sample tasks to everyone."),

        ("41-task-add", "Add and complete a task", [
            "Add a task with the plus button.",
            "Tap the task to mark it done.",
            "Close the app and reopen it.",
        ], "The task is saved, and completing it sticks after a restart.",
         "The Add button used to open a message saying nothing, and ticking "
           "a task was forgotten on restart."),
    ]),

    ("Reports and the rest", [
        ("42-reports", "Reports", [
            "Open Reports after running a pump a few times.",
        ], "Water used, hours run and electricity, by day and week, with a "
           "comparison against the period before. Soil rated in plain words.",
         "Harvest performance is deliberately empty. Nothing in the app "
           "records a sowing or a harvest, so there is no yield to compare. "
           "It used to show 91.7% efficiency to everyone."),

        ("43-analytics", "Farm analytics", ["Open Farm Analytics."],
         "A health score per field and plain observations drawn from your "
           "readings.",
         "Satellite crop health and yield prediction are empty on purpose. "
           "The first needs a satellite imagery provider, the second needs a "
           "season of harvest records. Both used to show invented figures."),

        ("44-schemes", "Government schemes", ["Open Government Schemes."],
         "PM-KISAN, Fasal Bima, KCC, Soil Health Card, eNAM and others, with "
           "eligibility and how to apply.", None),

        ("45-profile", "My profile", [
            "Settings, then My Profile.",
            "Change the farm name and save.",
            "Close the app and reopen it.",
        ], "The farm name and location are filled in from setup, and the "
           "change survives a restart.",
         "This screen used to open blank and save nothing."),

        ("46-settings", "Settings and language", [
            "Change the language to Hindi or Punjabi.",
        ], "The app changes language immediately and stays that way after a "
           "restart.",
         "10 languages: English, Hindi, Punjabi, Marathi, Telugu, Tamil, "
           "Kannada, Bengali, Gujarati and Malayalam."),
    ]),
]

KNOWN_LIMITS = [
    ("Sign in by phone (OTP)",
     "Built but cannot deliver. No SMS reaches an Indian number until DLT "
     "registration with TRAI is complete, or the company's existing SMS "
     "account is shared with us. Email sign-in works and is the way in for "
     "now."),
    ("Sign in by username",
     "There is no username in the app - accounts are identified by email or "
     "phone. The tab that used to fail when tapped has been removed and "
     "replaced with code sign-in. Not a defect."),
    ("Marketplace and mandi prices",
     "The screens work but nothing is stored, so listings are not saved and "
     "every tester sees the same sample data. Not yet built, not broken."),
    ("The first request of the day is slow",
     "Up to a minute. The server sleeps when unused. A spinner shows "
     "throughout. This is the hosting plan."),
    ("The AI field monitor needs the APK",
     "It cannot run in Expo Go. Test it only on the installed build."),
    ("No physical sensors exist",
     "Soil and device readings are tested with simulated data. Blank soil "
     "readings on a new account are correct."),
]


def placeholder(doc, name):
    """A visible box naming the screenshot that belongs here."""
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.width = Inches(2.6)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('[ screenshot: %s.png ]' % name)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run('capture with scripts/capture-screens.sh')
    r.italic = True
    r.font.size = Pt(7.5)
    r.font.color.rgb = GREY


def build():
    doc = Document()
    for s in doc.sections:
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10.5)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run('SmartKisan - Testing Manual')
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = GREEN

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run('Every screen in the Android app, what to do on it, and what '
                  'should happen')
    r.font.size = Pt(11)
    r.font.color.rgb = GREY

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run('Hbeonlabs Technologies Pvt. Ltd.  |  27 September 2026  |  '
                  'app-release.apk')
    r.font.size = Pt(9)
    r.font.color.rgb = GREY

    # ── Read this first ──────────────────────────────────────────────────
    doc.add_heading('Before you start', level=1)
    doc.add_paragraph(
        'Please read the known limits at the end of this document first. '
        'Several things in the app will look wrong and are not - a scan that '
        'refuses to answer, a blank soil reading, an empty harvest figure. '
        'Each is deliberate and the reason is given. Knowing them in advance '
        'saves filing defects that are not defects.')
    doc.add_paragraph(
        'Where a step mentions something that used to be broken, it is worth '
        'testing closely - those are the repairs from the last round and the '
        'most likely place for something to have been missed.')

    total = sum(len(items) for _, items in SECTIONS)
    have = len([f for f in os.listdir(SHOTS)]) if os.path.isdir(SHOTS) else 0
    p = doc.add_paragraph()
    r = p.add_run('%d screens to test.' % total)
    r.bold = True
    if have == 0:
        p.add_run(' Screenshots have not been captured yet - each one shows a '
                  'box naming the image that belongs there. Run '
                  'scripts/capture-screens.sh with a phone connected, then '
                  'generate this document again.')

    # ── The walkthrough ──────────────────────────────────────────────────
    n = 0
    for heading, items in SECTIONS:
        doc.add_page_break()
        doc.add_heading(heading, level=1)

        for shot, title, steps, expected, note in items:
            n += 1
            h = doc.add_heading(level=2)
            r = h.add_run('%d. %s' % (n, title))
            r.font.color.rgb = GREEN

            for step in steps:
                p = doc.add_paragraph(step, style='List Bullet')
                p.paragraph_format.space_after = Pt(2)

            p = doc.add_paragraph()
            r = p.add_run('Expected: ')
            r.bold = True
            r.font.color.rgb = DARK
            p.add_run(expected)

            if note:
                p = doc.add_paragraph()
                r = p.add_run('Please note: ')
                r.bold = True
                r.font.color.rgb = AMBER
                r = p.add_run(note)
                r.font.color.rgb = DARK

            path = os.path.join(SHOTS, shot + '.png')
            if os.path.exists(path):
                try:
                    doc.add_picture(path, width=Inches(2.6))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    placeholder(doc, shot)
            else:
                placeholder(doc, shot)

            doc.add_paragraph()

    # ── Known limits ─────────────────────────────────────────────────────
    doc.add_page_break()
    doc.add_heading('Known limits - please do not file these as defects',
                    level=1)
    doc.add_paragraph(
        'Each of these is understood and either deliberate or waiting on '
        'something outside development.')

    for name, text in KNOWN_LIMITS:
        p = doc.add_paragraph()
        r = p.add_run(name)
        r.bold = True
        r.font.color.rgb = AMBER
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(10)

    doc.add_heading('Reporting a defect', level=1)
    doc.add_paragraph(
        'So that a report can be acted on without going back and forth, '
        'please include:')
    for line in [
        'The screen number from this document.',
        'What you did, what happened, and what you expected.',
        'The exact words of any message on screen - a screenshot is best. '
        'This is the single most useful thing: a defect last round could not '
        'be reproduced for two weeks because the error text was not recorded.',
        'Whether it happened once or every time.',
        'The phone model and Android version.',
    ]:
        doc.add_paragraph(line, style='List Bullet')

    out = os.path.join(HERE, FILENAME)
    doc.save(out)
    print('Wrote ' + out)
    print('  %d screens, %d screenshots found in docs/screenshots'
          % (total, have))
    if have == 0:
        print('  Capture them:  bash scripts/capture-screens.sh')


if __name__ == '__main__':
    build()
